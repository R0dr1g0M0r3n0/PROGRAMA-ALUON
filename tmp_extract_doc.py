from docx import Document
import sys
p=sys.argv[1]
doc=Document(p)
for i, para in enumerate(doc.paragraphs):
    text=para.text.strip()
    if text:
        print(f"P{i}: {text}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}")
    for row in table.rows:
        print(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
